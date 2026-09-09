"""
Parser สำหรับไฟล์ export รูปแบบใหม่ ("สำเนาของข้อมูลดิบ") — พบจากไฟล์ตัวอย่างจริงที่ผู้ใช้ส่งมา
(Cancelled.xlsx / Calendar.xlsx ใน Downloads) คนละโครงสร้างกับไฟล์ตัวอย่างชุดแรก (Cancelled.xlsx /
Function_Calendar.xlsx ที่ root โปรเจกต์ ซึ่ง parse ด้วย app/parsing/cancelled.py + calendar.py เดิม)

โครงสร้างจริงที่พบ — ตารางแบน 1 แถวต่อ 1 รายการ, header แถวที่ 1:
  วัน | dd/mm/yy | Status | Cust Type | Time | Event Type | location | Customer | Note | Sales | Pax

ข้อดี: Cust Type และ Pax เป็นคอลัมน์แยกอยู่แล้ว ไม่ต้องเดาจากข้อความเหมือนไฟล์ชุดแรก
ข้อควรระวัง:
  - Status สะกดต่างจากที่โค้ดเดิมคาด: "Cancelled" (ไม่ใช่ "Cancel") และ "Comfirmed" (สะกดผิดในไฟล์จริง
    ไม่ใช่ "Confirmed") — normalize ไว้ใน STATUS_MAP ด้านล่าง
  - งาน Cancelled ยังแกะรหัสประเภทงาน (MT/WD/DN/...) ได้จากข้อความในคอลัมน์ Note เหมือนไฟล์ชุดแรก
    (ใช้ find_job_type ตัวเดียวกัน) เพราะฟอร์แมตข้อความคล้ายเดิม (มี "[QN...] CXL. ... MT ... 50P")
  - งาน Confirmed/Pending **ไม่มีรหัสประเภทงานเลย** — คอลัมน์ Event Type เป็นข้อความกิจกรรมอิสระ
    (เช่น "งานประชุม", "อาหารกลางวัน") ไม่ตรงกับชุดรหัส MT/WD/DN — ผู้ใช้ยืนยันให้ปล่อยเป็น "ไม่ระบุ"
    ไปก่อน (จะกลับมาปรับ mapping ทีหลัง) ดู guess_job_type_from_event_type() ด้านล่าง: ตอนจะเปลี่ยน
    ทีหลัง แก้ logic แค่ในฟังก์ชันนี้ที่เดียว ไม่ต้องแตะที่อื่น
"""

import re
from datetime import date, datetime
from pathlib import Path

import openpyxl

from .cancelled import CUSTOMER_RE, categorize_reason, find_job_type

STATUS_MAP = {
    "cancelled": "Cancelled",
    "cancel": "Cancelled",
    "comfirmed": "Confirmed",  # สะกดผิดในไฟล์จริง
    "confirmed": "Confirmed",
    "tentative": "Pending",
    "not confirm": "Pending",
    "cut off": "Pending",
    "pending": "Pending",
}

VALID_CUSTOMER_TYPES = {"A", "B", "C", "N"}

DATE_CELL_RE = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{4})")
QN_RE = re.compile(r"\[?\s*(QN[\w-]+)\s*\]?")
TIME_RE = re.compile(r"(\d{1,2}):(\d{2})")


def is_raw_export_format(path: Path) -> bool:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [ws.cell(row=1, column=c).value for c in range(1, 4)]
    wb.close()
    return header == ["วัน", "dd/mm/yy", "Status"]


def _parse_date_cell(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        m = DATE_CELL_RE.search(value)
        if m:
            d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
            try:
                return date(y, mo, d)
            except ValueError:
                return None
    return None


def _parse_pax_cell(value) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        return None


def _time_of_day(time_text) -> str:
    if isinstance(time_text, str):
        m = TIME_RE.search(time_text)
        if m:
            hour = int(m.group(1))
            if hour < 12:
                return "ช่วงเช้า"
            if hour < 17:
                return "ช่วงบ่าย"
            return "ช่วงเย็น"
    return "ช่วงเช้า"


def guess_job_type_from_event_type(event_type_text: str | None) -> str | None:
    """
    Placeholder — ยังไม่มี mapping ระหว่างข้อความกิจกรรมอิสระ (Event Type) กับรหัสประเภทงาน
    (MT/WD/DN/...) จนกว่าทีมขายจะให้ mapping จริงมา ตอนนี้คืน None เสมอ (= "ไม่ระบุ" ในกราฟ)
    แก้ทีหลัง: เติม logic ในฟังก์ชันนี้ที่เดียว (เช่น dict คำ -> รหัส) ไม่ต้องแก้ที่อื่นในโปรเจกต์
    """
    return None


def _read_raw_rows(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = []
    for r in range(2, ws.max_row + 1):
        status_raw = ws.cell(row=r, column=3).value
        date_val = ws.cell(row=r, column=2).value
        if not status_raw or not date_val:
            continue  # แถวว่างท้ายชีท (พบเยอะในไฟล์ export จริง)

        status = STATUS_MAP.get(str(status_raw).strip().lower(), str(status_raw).strip())
        event_date = _parse_date_cell(date_val)

        customer_type = ws.cell(row=r, column=4).value
        customer_type = str(customer_type).strip().upper() if customer_type else None
        if customer_type not in VALID_CUSTOMER_TYPES:
            customer_type = None

        time_text = ws.cell(row=r, column=5).value
        event_type_text = ws.cell(row=r, column=6).value
        event_type_text = str(event_type_text).strip() if event_type_text else None
        customer_name = ws.cell(row=r, column=8).value
        customer_name = str(customer_name).strip() if customer_name else None
        note = ws.cell(row=r, column=9).value
        note = str(note).strip() if note else ""
        sales = ws.cell(row=r, column=10).value
        sales = str(sales).strip() if sales else None
        pax = _parse_pax_cell(ws.cell(row=r, column=11).value)

        if not customer_type and note:
            m = CUSTOMER_RE.search(note)
            if m:
                customer_type = m.group(1).upper()

        rows.append({
            "status": status,
            "event_date": event_date,
            "customer_type": customer_type,
            "time_of_day": _time_of_day(time_text),
            "event_type_text": event_type_text,
            "customer_name": customer_name,
            "note": note,
            "sales": sales,
            "pax": pax,
        })
    wb.close()
    return rows


def load_cancelled_rows(path: Path) -> list[dict]:
    """คืนแถวสถานะ Cancelled เท่านั้น รูปแบบ dict เดียวกับ app/parsing/cancelled.py::load_cancelled_rows"""
    out = []
    for row in _read_raw_rows(path):
        if row["status"] != "Cancelled":
            continue
        job_type, pos = find_job_type(row["note"]) if row["note"] else (None, -1)
        raw_reason = row["note"][:pos] if pos >= 0 else row["note"]
        raw_reason = re.sub(r"^\[?QN[\w-]*\]?\s*", "", raw_reason)
        raw_reason = re.sub(r"^CX[LK]\.?", "", raw_reason).strip()
        qn_match = QN_RE.search(row["note"])
        out.append({
            "qtn": qn_match.group(1) if qn_match else None,
            "event_date": row["event_date"],
            "customer_name": row["customer_name"],
            "sales": row["sales"],
            "job_type": job_type,
            "customer_type": row["customer_type"],
            "pax": row["pax"],
            "raw_reason": raw_reason,
            "reason_category": categorize_reason(raw_reason),
        })
    return out


def load_calendar_events(path: Path) -> list[dict]:
    """คืนแถวสถานะไม่ใช่ Cancelled (Confirmed/Pending) รูปแบบ dict เดียวกับ app/parsing/calendar.py::extract_events"""
    out = []
    for row in _read_raw_rows(path):
        if row["status"] == "Cancelled":
            continue
        title = " / ".join(p for p in [row["customer_name"], row["event_type_text"]] if p) or "(ไม่มีชื่องาน)"
        out.append({
            "date": row["event_date"].isoformat() if row["event_date"] else None,
            "date_obj": row["event_date"],
            "start": None,
            "end": None,
            "time_of_day": row["time_of_day"],
            "pax": row["pax"] or 0,
            "sales": row["sales"],
            "title": title,
            "job_type": guess_job_type_from_event_type(row["event_type_text"]),
            "status": row["status"],
        })
    return out
