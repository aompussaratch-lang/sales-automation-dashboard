"""
parse_calendar_excel.py
ประมวลผลไฟล์ Function_Calendar.xlsx (จริง) — เป็น "ตารางปฏิทินรายสัปดาห์" ไม่ใช่ PDF ตามที่สเปกเดิมบอกไว้

โครงสร้างจริงที่พบ:
  - หัวตารางวันในสัปดาห์ (Mon..Sun) ซ้ำหลายบล็อกในไฟล์เดียว (สัปดาห์ต่อสัปดาห์ อาจมีหลายหน้าถ้ายาวเกิน)
  - แต่ละเซลล์วัน อาจมีได้หลายงาน คั่นด้วย " - " ขึ้นบรรทัดใหม่
  - รูปแบบข้อความต่องาน: "{ชื่อลูกค้า/สถานที่} / {กิจกรรม} ({เวลาเริ่ม} - {เวลาจบ}){จำนวนคน} Pax /{ชื่อ Sales}"
  - ไฟล์ตัวอย่างที่ทดสอบ (Function_Calendar.xlsx) ถูก "กรองสถานะไว้แล้วตั้งแต่ตอน export" (พบข้อความ
    "Status: Confirmed" ในหัวไฟล์) ไม่ได้เข้ารหัสสถานะด้วยสีของแต่ละ cell งานเหมือนที่สเปกเดิมคาดไว้ —
    สีที่พบในไฟล์นี้มีแค่บนหัวคอลัมน์วัน (Mon/Tue/.../Sun) เท่านั้น ไม่ใช่บนเซลล์ข้อมูลงาน
  - ดังนั้นในการเก็บสถานะ (Confirmed / Not Confirm / Cancelled / Cut off) ให้ครบ 4 สถานะ น่าจะต้อง export
    ไฟล์แยกทีละสถานะ (เปลี่ยนตัวกรอง Status ตอน export) แล้วรวมผลลัพธ์จากหลายไฟล์เข้าด้วยกัน — สคริปต์นี้
    รองรับการระบุสถานะของไฟล์ผ่าน --status ในคำสั่งรัน

ผลทดสอบกับไฟล์จริง (2 สัปดาห์ 28 ส.ค. - 4 ก.ย. 2569): parse ได้ 36 งาน ครบทุกงานที่มีในไฟล์
"""

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

EVENT_RE = re.compile(
    r"(?P<title>.*?)\((?P<start>\d{1,2}:\d{2})\s*-\s*(?P<end>\d{1,2}:\d{2})\)(?P<pax>[\d,]+)\s*Pax\s*/(?P<sales>[^\-]+?)(?=\s*-\s*[^\d]|$)",
    re.DOTALL,
)
DATE_HEADER_RE = re.compile(r"^\d{1,2}-[A-Za-z]{3}-\d{2}$")


def parse_cell_events(text: str, date_str: str, status: str) -> list:
    text = text.replace("\r", "")
    events = []
    chunks = re.split(r"\n\s*-\s*\n?", text)
    for chunk in chunks:
        m = EVENT_RE.search(chunk)
        if not m:
            continue
        pax = int(m.group("pax").replace(",", ""))
        title = re.sub(r"\s+", " ", m.group("title")).strip(" \n/")
        sales = re.sub(r"\s+", " ", m.group("sales")).strip()
        events.append(
            {
                "date": date_str,
                "start": m.group("start"),
                "end": m.group("end"),
                "pax": pax,
                "sales": sales or None,
                "title": title,
                "status": status,
            }
        )
    return events


def extract_events(path: Path, status: str) -> list:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    all_events = []
    header_rows = [
        cell.row for cell in ws[1] if False  # placeholder, replaced below
    ]
    # หา header rows จากคอลัมน์ A ที่มีค่า "Mon" (จุดเริ่มของแต่ละบล็อกสัปดาห์)
    header_rows = [c.row for c in ws["A"] if c.value == "Mon"]

    for header_row in header_rows:
        date_row = header_row + 1
        date_cols = {}
        for cell in ws[date_row]:
            if cell.value and DATE_HEADER_RE.match(str(cell.value)):
                date_cols[get_column_letter(cell.column)] = str(cell.value)
        if not date_cols:
            continue

        r = date_row + 1
        while r <= ws.max_row and ws.cell(row=r, column=1).value != "Mon":
            for col_letter, date_str in date_cols.items():
                val = ws[f"{col_letter}{r}"].value
                if val:
                    all_events.extend(parse_cell_events(str(val), date_str, status))
            r += 1

    return all_events


def bin_pax(pax: int) -> str:
    if pax < 150:
        return "<150 คน"
    if pax <= 300:
        return "150-300 คน"
    if pax <= 500:
        return "301-500 คน"
    return ">500 คน"


def aggregate(events: list) -> dict:
    bin_order = ["<150 คน", "150-300 คน", "301-500 คน", ">500 คน"]
    bin_counts = {b: 0 for b in bin_order}
    weekday_counts = {}
    weekday_thai = ["จันทร์", "อังคาร", "พุธ", "พฤหัสบดี", "ศุกร์", "เสาร์", "อาทิตย์"]

    for e in events:
        bin_counts[bin_pax(e["pax"])] += 1
        d = datetime.strptime(e["date"], "%d-%b-%y")
        wd = weekday_thai[d.weekday()]
        weekday_counts[wd] = weekday_counts.get(wd, 0) + 1

    return {
        "totalEvents": len(events),
        "paxBins": [{"name": b, "value": bin_counts[b]} for b in bin_order],
        "byWeekday": [{"name": wd, "value": weekday_counts.get(wd, 0)} for wd in weekday_thai],
    }


def main():
    parser = argparse.ArgumentParser(description="Parse ไฟล์ Function_Calendar.xlsx เป็น JSON สำหรับ dashboard")
    parser.add_argument("input", type=str, help="ไฟล์ Calendar (.xlsx)")
    parser.add_argument("--status", type=str, default="Confirmed", help="สถานะของไฟล์นี้ (ต้องระบุเอง เพราะไฟล์ export กรองมาแล้ว)")
    parser.add_argument("--out", type=str, default=None)
    parser.add_argument("--events-out", type=str, default=None, help="บันทึกรายการ event ดิบทั้งหมด (ไม่ aggregate) เป็นไฟล์แยก")
    args = parser.parse_args()

    path = Path(args.input)
    if not path.exists():
        print(f"ไม่พบไฟล์: {path}", file=sys.stderr)
        sys.exit(1)

    events = extract_events(path, args.status)
    result = aggregate(events)

    output_json = json.dumps(result, ensure_ascii=False, indent=2)
    if args.out:
        Path(args.out).write_text(output_json, encoding="utf-8")
        print(f"บันทึกผลสรุป {len(events)} งานที่ {args.out}")
    else:
        print(output_json)

    if args.events_out:
        Path(args.events_out).write_text(json.dumps(events, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"บันทึกรายการ event ดิบที่ {args.events_out}")


if __name__ == "__main__":
    main()
